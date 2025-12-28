"""
Tamil Text Spacing Corrector for Quran Excel Files - V2
========================================================
This application fixes common Tamil spacing issues:
1. Removes spaces before Tamil vowel signs (matras) - they must attach to consonants
2. Removes spaces before pulli (virama - ்)
3. Normalizes multiple spaces to single space
4. Trims leading/trailing whitespace

Author: Created for Spiral Nineteen IT Technologies
"""

import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import traceback

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from pathlib import Path


# Tamil Unicode ranges and characters
TAMIL_VOWEL_SIGNS = '\u0BBE\u0BBF\u0BC0\u0BC1\u0BC2\u0BC6\u0BC7\u0BC8\u0BCA\u0BCB\u0BCC'  # ா ி ீ ு ூ ெ ே ை ொ ோ ௌ
TAMIL_PULLI = '\u0BCD'  # ் (virama)
TAMIL_ANUSVARA = '\u0B82'  # ஂ
TAMIL_VISARGA = '\u0B83'  # ஃ

# Combined pattern for all characters that should not have space before them
TAMIL_DEPENDENT_MARKS = TAMIL_VOWEL_SIGNS + TAMIL_PULLI + TAMIL_ANUSVARA


def fix_tamil_spacing(text):
    """
    Fix spacing issues in Tamil text.
    """
    if not isinstance(text, str):
        return text
    
    if not text.strip():
        return text
    
    original = text
    
    # Pattern 1: Remove space(s) before Tamil dependent vowel signs and pulli
    pattern_dependent = r'\s+([' + TAMIL_DEPENDENT_MARKS + '])'
    text = re.sub(pattern_dependent, r'\1', text)
    
    # Pattern 2: Remove zero-width characters
    text = re.sub(r'[\u200B\u200C\u200D\uFEFF]+', '', text)
    
    # Pattern 3: Normalize multiple spaces to single space
    text = re.sub(r' {2,}', ' ', text)
    
    # Pattern 4: Trim leading and trailing whitespace
    text = text.strip()
    
    return text


def detect_tamil_issues(text):
    """
    Detect and report Tamil spacing issues in text.
    Returns a list of issues found.
    """
    issues = []
    if not isinstance(text, str):
        return issues
    
    # Check for space before dependent marks
    pattern = r'.{0,5}\s+[' + TAMIL_DEPENDENT_MARKS + '].{0,5}'
    matches = re.findall(pattern, text)
    for match in matches:
        issues.append(f"Space before matra: ...{match.strip()}...")
    
    # Check for multiple spaces
    if re.search(r' {2,}', text):
        issues.append("Multiple consecutive spaces")
    
    # Check for zero-width characters
    if re.search(r'[\u200B\u200C\u200D\uFEFF]', text):
        issues.append("Zero-width characters found")
    
    return issues


def has_tamil_text(text):
    """Check if text contains Tamil characters."""
    if not isinstance(text, str):
        return False
    return bool(re.search(r'[\u0B80-\u0BFF]', text))


class TamilFixerApp:
    """GUI Application for Tamil Text Fixer"""
    
    def __init__(self, root):
        self.root = root
        self.root.title("Tamil Quran Text Corrector - Spiral Nineteen")
        self.root.geometry("750x550")
        self.root.resizable(True, True)
        
        self.setup_ui()
        self.input_file = None
        
        # Check dependencies
        self.check_dependencies()
    
    def check_dependencies(self):
        """Check if required libraries are installed."""
        missing = []
        if not PANDAS_AVAILABLE:
            missing.append("pandas")
        if not OPENPYXL_AVAILABLE:
            missing.append("openpyxl")
        
        if missing:
            msg = f"Missing libraries: {', '.join(missing)}\n\nPlease install using:\npip install {' '.join(missing)}"
            messagebox.showerror("Missing Dependencies", msg)
    
    def setup_ui(self):
        # Main frame
        main_frame = ttk.Frame(self.root, padding="20")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Title
        title_label = ttk.Label(
            main_frame, 
            text="Tamil Quran Text Corrector",
            font=('Helvetica', 14, 'bold')
        )
        title_label.pack(pady=(0, 5))
        
        subtitle_label = ttk.Label(
            main_frame,
            text="Bismillah hir Rahman nir Raheem",
            font=('Arial', 10, 'italic')
        )
        subtitle_label.pack(pady=(0, 15))
        
        # Description
        desc_text = """This tool fixes common Tamil text spacing issues:
• Removes unwanted spaces before vowel signs (matras)
• Fixes broken character combinations  
• Normalizes multiple spaces
• Removes zero-width characters"""
        
        desc_label = ttk.Label(main_frame, text=desc_text, justify=tk.LEFT)
        desc_label.pack(pady=(0, 15), anchor=tk.W)
        
        # File selection frame
        file_frame = ttk.LabelFrame(main_frame, text="Select Excel File", padding="10")
        file_frame.pack(fill=tk.X, pady=(0, 10))
        
        self.file_label = ttk.Label(file_frame, text="No file selected")
        self.file_label.pack(side=tk.LEFT, fill=tk.X, expand=True)
        
        browse_btn = ttk.Button(file_frame, text="Browse...", command=self.browse_file)
        browse_btn.pack(side=tk.RIGHT, padx=(10, 0))
        
        # Options frame
        options_frame = ttk.LabelFrame(main_frame, text="Options", padding="10")
        options_frame.pack(fill=tk.X, pady=(0, 10))
        
        ttk.Label(options_frame, text="Tamil Column (leave empty for auto-detect):").pack(anchor=tk.W)
        self.column_entry = ttk.Entry(options_frame, width=40)
        self.column_entry.pack(anchor=tk.W, pady=(5, 0))
        
        # Progress frame
        progress_frame = ttk.Frame(main_frame)
        progress_frame.pack(fill=tk.X, pady=(0, 10))
        
        self.progress_label = ttk.Label(progress_frame, text="Ready", foreground="blue")
        self.progress_label.pack(anchor=tk.W)
        
        self.progress_bar = ttk.Progressbar(progress_frame, mode='determinate', length=400)
        self.progress_bar.pack(fill=tk.X, pady=(5, 0))
        
        # Buttons frame
        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill=tk.X, pady=(5, 0))
        
        self.process_btn = ttk.Button(
            btn_frame, 
            text="Process & Fix Tamil Text",
            command=self.process_file
        )
        self.process_btn.pack(side=tk.LEFT)
        
        self.preview_btn = ttk.Button(
            btn_frame,
            text="Preview Issues",
            command=self.preview_issues
        )
        self.preview_btn.pack(side=tk.LEFT, padx=(10, 0))
        
        clear_btn = ttk.Button(
            btn_frame,
            text="Clear Log",
            command=self.clear_log
        )
        clear_btn.pack(side=tk.RIGHT)
        
        # Results text area
        results_frame = ttk.LabelFrame(main_frame, text="Results / Log", padding="10")
        results_frame.pack(fill=tk.BOTH, expand=True, pady=(10, 0))
        
        # Create text widget with scrollbar
        text_frame = ttk.Frame(results_frame)
        text_frame.pack(fill=tk.BOTH, expand=True)
        
        self.results_text = tk.Text(text_frame, height=10, wrap=tk.WORD, font=('Consolas', 9))
        scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.results_text.yview)
        
        self.results_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.results_text.config(yscrollcommand=scrollbar.set)
        
        # Initial message
        self.log("Application ready. Select an Excel file to begin.\n")
    
    def clear_log(self):
        """Clear the results text area."""
        self.results_text.delete(1.0, tk.END)
    
    def browse_file(self):
        """Browse for Excel file."""
        file_path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[
                ("Excel Files", "*.xlsx *.xls"),
                ("All Files", "*.*")
            ]
        )
        if file_path:
            self.input_file = file_path
            filename = Path(file_path).name
            self.file_label.config(text=filename)
            self.log(f"\n✓ Selected: {filename}\n")
            self.log(f"  Full path: {file_path}\n")
    
    def update_progress(self, message, value):
        """Update progress bar and label."""
        self.progress_label.config(text=message)
        self.progress_bar['value'] = value
        self.root.update_idletasks()
    
    def log(self, message):
        """Log message to results text area."""
        self.results_text.insert(tk.END, message)
        self.results_text.see(tk.END)
        self.root.update_idletasks()
    
    def preview_issues(self):
        """Preview Tamil spacing issues in the file."""
        if not self.input_file:
            messagebox.showwarning("No File", "Please select an Excel file first.")
            return
        
        self.log("\n" + "="*50 + "\n")
        self.log("SCANNING FOR TAMIL SPACING ISSUES...\n")
        self.log("="*50 + "\n\n")
        
        self.update_progress("Reading file...", 10)
        self.preview_btn.config(state=tk.DISABLED)
        
        try:
            # Read Excel file
            self.log(f"Reading file: {Path(self.input_file).name}\n")
            df = pd.read_excel(self.input_file)
            self.log(f"✓ Loaded {len(df)} rows, {len(df.columns)} columns\n\n")
            
            self.update_progress("Detecting Tamil columns...", 20)
            
            # Show all columns
            self.log(f"Columns in file: {list(df.columns)}\n\n")
            
            # Detect Tamil columns
            tamil_columns = []
            for col in df.columns:
                col_str = str(col).lower()
                if 'tamil' in col_str:
                    tamil_columns.append(col)
                    self.log(f"✓ Found Tamil column by name: '{col}'\n")
                else:
                    # Check content for Tamil characters
                    sample = df[col].dropna().head(20)
                    tamil_count = sum(1 for val in sample if has_tamil_text(val))
                    if tamil_count > 0:
                        tamil_columns.append(col)
                        self.log(f"✓ Found Tamil column by content: '{col}' ({tamil_count} Tamil texts in sample)\n")
            
            if not tamil_columns:
                self.log("\n❌ No Tamil columns detected in the file!\n")
                self.log("Please check if the file contains Tamil text.\n")
                self.update_progress("No Tamil columns found", 100)
                return
            
            self.log(f"\nTotal Tamil columns: {len(tamil_columns)}\n")
            self.log("-"*50 + "\n\n")
            
            # Scan for issues
            self.update_progress("Scanning for issues...", 40)
            
            issue_count = 0
            rows_with_issues = 0
            examples_shown = 0
            max_examples = 15
            
            for col in tamil_columns:
                self.log(f"Scanning column: '{col}'...\n")
                col_issues = 0
                
                for idx in range(len(df)):
                    val = df.at[idx, col]
                    if isinstance(val, str):
                        issues = detect_tamil_issues(val)
                        if issues:
                            col_issues += 1
                            rows_with_issues += 1
                            
                            # Show some examples
                            if examples_shown < max_examples:
                                row_num = idx + 2  # Excel row (1-indexed + header)
                                self.log(f"  Row {row_num}: {issues[0]}\n")
                                
                                # Show before/after preview
                                original = val[:60] + "..." if len(val) > 60 else val
                                fixed = fix_tamil_spacing(val)
                                fixed_preview = fixed[:60] + "..." if len(fixed) > 60 else fixed
                                
                                if val != fixed:
                                    self.log(f"    Before: {original}\n")
                                    self.log(f"    After:  {fixed_preview}\n")
                                
                                examples_shown += 1
                    
                    # Update progress
                    if idx % 500 == 0:
                        progress = 40 + int((idx / len(df)) * 50)
                        self.update_progress(f"Scanning row {idx}/{len(df)}...", progress)
                
                self.log(f"  → Found {col_issues} rows with issues in '{col}'\n\n")
                issue_count += col_issues
            
            # Summary
            self.log("="*50 + "\n")
            self.log("SCAN SUMMARY\n")
            self.log("="*50 + "\n")
            self.log(f"Total rows scanned: {len(df)}\n")
            self.log(f"Rows with issues: {rows_with_issues}\n")
            self.log(f"Tamil columns: {', '.join(str(c) for c in tamil_columns)}\n")
            
            if rows_with_issues > 0:
                self.log(f"\n→ Click 'Process & Fix Tamil Text' to fix these issues.\n")
            else:
                self.log(f"\n✓ No spacing issues found! The Tamil text looks good.\n")
            
            self.update_progress("Scan complete", 100)
            
        except Exception as e:
            error_msg = f"Error: {str(e)}\n\nDetails:\n{traceback.format_exc()}"
            self.log(f"\n❌ ERROR:\n{error_msg}\n")
            messagebox.showerror("Error", f"Error scanning file:\n{str(e)}")
        
        finally:
            self.preview_btn.config(state=tk.NORMAL)
    
    def process_file(self):
        """Process and fix Tamil text in the file."""
        if not self.input_file:
            messagebox.showwarning("No File", "Please select an Excel file first.")
            return
        
        self.log("\n" + "="*50 + "\n")
        self.log("PROCESSING FILE...\n")
        self.log("="*50 + "\n\n")
        
        self.update_progress("Reading file...", 5)
        self.process_btn.config(state=tk.DISABLED)
        self.preview_btn.config(state=tk.DISABLED)
        
        try:
            # Read Excel file
            df = pd.read_excel(self.input_file)
            self.log(f"✓ Loaded {len(df)} rows\n")
            
            self.update_progress("Detecting Tamil columns...", 10)
            
            # Detect Tamil columns
            tamil_columns = []
            specified_col = self.column_entry.get().strip()
            
            if specified_col:
                if specified_col in df.columns:
                    tamil_columns = [specified_col]
                    self.log(f"✓ Using specified column: '{specified_col}'\n")
                else:
                    self.log(f"❌ Column '{specified_col}' not found!\n")
                    self.log(f"   Available columns: {list(df.columns)}\n")
                    messagebox.showerror("Error", f"Column '{specified_col}' not found in file.")
                    return
            else:
                # Auto-detect
                for col in df.columns:
                    col_str = str(col).lower()
                    if 'tamil' in col_str:
                        tamil_columns.append(col)
                    else:
                        sample = df[col].dropna().head(20)
                        if any(has_tamil_text(val) for val in sample):
                            tamil_columns.append(col)
            
            if not tamil_columns:
                self.log("❌ No Tamil columns detected!\n")
                messagebox.showerror("Error", "No Tamil columns found in the file.")
                return
            
            self.log(f"✓ Tamil columns to process: {tamil_columns}\n\n")
            
            # Process each column
            total_fixes = 0
            
            for col in tamil_columns:
                self.log(f"Processing column: '{col}'...\n")
                fixes = 0
                
                for idx in range(len(df)):
                    original = df.at[idx, col]
                    if isinstance(original, str):
                        fixed = fix_tamil_spacing(original)
                        if fixed != original:
                            df.at[idx, col] = fixed
                            fixes += 1
                    
                    # Update progress
                    if idx % 200 == 0:
                        progress = 20 + int((idx / len(df)) * 60)
                        self.update_progress(f"Processing row {idx + 1}/{len(df)}...", progress)
                
                self.log(f"  → Fixed {fixes} rows in '{col}'\n")
                total_fixes += fixes
            
            # Generate output path
            input_path = Path(self.input_file)
            output_path = input_path.parent / f"{input_path.stem}_corrected{input_path.suffix}"
            
            # Save file
            self.update_progress("Saving corrected file...", 90)
            self.log(f"\nSaving to: {output_path.name}...\n")
            
            df.to_excel(output_path, index=False)
            
            self.update_progress("Complete!", 100)
            
            # Summary
            self.log("\n" + "="*50 + "\n")
            self.log("✓ PROCESSING COMPLETE - Alhamdulillah!\n")
            self.log("="*50 + "\n")
            self.log(f"Total rows processed: {len(df)}\n")
            self.log(f"Total fixes applied: {total_fixes}\n")
            self.log(f"Output file: {output_path}\n")
            
            messagebox.showinfo(
                "Success - Alhamdulillah!",
                f"File processed successfully!\n\n"
                f"Rows fixed: {total_fixes}\n\n"
                f"Saved to:\n{output_path.name}"
            )
            
        except Exception as e:
            error_msg = f"Error: {str(e)}\n\nDetails:\n{traceback.format_exc()}"
            self.log(f"\n❌ ERROR:\n{error_msg}\n")
            messagebox.showerror("Error", f"Error processing file:\n{str(e)}")
        
        finally:
            self.process_btn.config(state=tk.NORMAL)
            self.preview_btn.config(state=tk.NORMAL)


def main():
    """Main entry point"""
    root = tk.Tk()
    app = TamilFixerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
